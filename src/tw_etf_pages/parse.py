"""Parse issuer Excel / third-party JSON into normalized snapshots."""

from __future__ import annotations

import hashlib
import logging
from datetime import date
from pathlib import Path
from typing import Any

import warnings

from openpyxl import load_workbook

from .config import AppConfig, PlaceholderPolicy
from .utils import (
    normalize_stock_code,
    now_taipei_iso,
    parse_int,
    parse_iso_or_slash_date,
    parse_roc_date,
    parse_weight_pct,
    read_json,
)

logger = logging.getLogger(__name__)

def _load_rows(path: Path) -> list[tuple]:
    with warnings.catch_warnings():
        warnings.filterwarnings(
            "ignore",
            message="Workbook contains no default style.*",
            category=UserWarning,
        )
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    return rows




def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def is_placeholder(
    shares: int,
    weight_pct: float,
    policy: PlaceholderPolicy,
) -> bool:
    return shares <= policy.max_shares and weight_pct < policy.max_weight_pct


def _holding(
    stock_code: str,
    stock_name: str,
    shares: int,
    weight_pct: float,
    policy: PlaceholderPolicy,
    market_value: float | None = None,
) -> dict[str, Any]:
    code = normalize_stock_code(stock_code)
    row: dict[str, Any] = {
        "stock_code": code,
        "stock_name": str(stock_name).strip(),
        "shares": shares,
        "weight_pct": round(weight_pct, 4),
        "is_placeholder": is_placeholder(shares, weight_pct, policy),
    }
    if market_value is not None:
        row["market_value"] = market_value
    return row


def parse_uni_excel(
    path: Path,
    ticker: str,
    policy: PlaceholderPolicy,
    source: str = "uni_excel",
) -> dict[str, Any]:
    """Parse Uni (ezmoney) AssetExcelNPOI workbook."""
    rows = _load_rows(path)

    as_of: date | None = None
    holdings: list[dict[str, Any]] = []
    in_stocks = False
    meta: dict[str, Any] = {}

    for row in rows:
        cell0 = row[0] if row else None
        if cell0 is None:
            if in_stocks:
                # blank row after stocks ends the table
                if holdings:
                    break
            continue
        text0 = str(cell0).strip()
        if text0.startswith("資料日期"):
            as_of = parse_roc_date(text0)
            continue
        if text0 == "淨資產" and len(row) > 1 and row[1] is not None:
            meta["nav_total_raw"] = str(row[1])
            continue
        if text0 == "流通在外單位數" and len(row) > 1 and row[1] is not None:
            meta["units_raw"] = str(row[1])
            continue
        if text0 == "每單位淨值" and len(row) > 1 and row[1] is not None:
            meta["nav_pu_raw"] = str(row[1])
            continue
        if text0 == "股票代號":
            in_stocks = True
            continue
        if in_stocks:
            # stop if we hit another section header
            if text0 in {"期貨代號", "項目", "期貨(名目本金)", "股票"}:
                break
            name = row[1] if len(row) > 1 else ""
            shares = parse_int(row[2] if len(row) > 2 else 0)
            weight = parse_weight_pct(row[3] if len(row) > 3 else 0)
            if not text0:
                continue
            holdings.append(_holding(text0, name or "", shares, weight, policy))

    if as_of is None:
        raise ValueError(f"{path}: missing 資料日期")
    if not holdings:
        raise ValueError(f"{path}: no stock holdings found")

    return _snapshot(
        ticker=ticker,
        as_of_date=as_of,
        holdings=holdings,
        source=source,
        raw_path=path,
        meta=meta,
    )


def parse_fh_excel(
    path: Path,
    ticker: str,
    policy: PlaceholderPolicy,
    source: str = "fh_excel",
) -> dict[str, Any]:
    """Parse FHTrust assetsExcel workbook."""
    rows = _load_rows(path)

    as_of: date | None = None
    holdings: list[dict[str, Any]] = []
    in_stocks = False
    meta: dict[str, Any] = {}

    for i, row in enumerate(rows):
        cell0 = row[0] if row else None
        if cell0 is None:
            continue
        text0 = str(cell0).strip()
        if text0.startswith("日期"):
            as_of = parse_iso_or_slash_date(text0)
            continue
        if text0 == "基金資產淨值" and i + 1 < len(rows) and rows[i + 1][0] is not None:
            meta["nav_total_raw"] = str(rows[i + 1][0])
            continue
        if text0 == "基金在外流通單位數" and i + 1 < len(rows) and rows[i + 1][0] is not None:
            meta["units_raw"] = str(rows[i + 1][0])
            continue
        if text0 == "基金每單位淨值" and i + 1 < len(rows) and rows[i + 1][0] is not None:
            meta["nav_pu_raw"] = str(rows[i + 1][0])
            continue
        if text0 == "證券代號":
            in_stocks = True
            continue
        if in_stocks:
            if not text0:
                break
            name = row[1] if len(row) > 1 else ""
            shares = parse_int(row[2] if len(row) > 2 else 0)
            mv = None
            if len(row) > 3 and row[3] is not None:
                try:
                    mv = float(parse_int(row[3]))
                except ValueError:
                    mv = None
            weight = parse_weight_pct(row[4] if len(row) > 4 else 0)
            holdings.append(
                _holding(text0, name or "", shares, weight, policy, market_value=mv)
            )

    if as_of is None:
        raise ValueError(f"{path}: missing 日期")
    if not holdings:
        raise ValueError(f"{path}: no stock holdings found")

    return _snapshot(
        ticker=ticker,
        as_of_date=as_of,
        holdings=holdings,
        source=source,
        raw_path=path,
        meta=meta,
    )


def parse_zdsetf_snapshot(
    path: Path,
    ticker: str,
    policy: PlaceholderPolicy,
) -> dict[str, Any]:
    """Fallback: zdsetf.com snapshot JSON.

    Field aliases (00405A / 00406A / 00980A / 00984A / 00400A / 00401A / 00987A / future issuers may differ slightly from 00992A):
    - date: snapshot_date | as_of_date | date
    - holdings list: holdings | positions | stocks
    - code/name: stock_code|code , stock_name|name
    - weight: weight_pct | weight | weight_percent (string or number OK)
    - market value: market_value | market_value_ntd | mv
    """
    data = read_json(path)
    raw_date = (
        data.get("snapshot_date")
        or data.get("as_of_date")
        or data.get("date")
    )
    if raw_date is None:
        raise ValueError(f"{path}: missing snapshot_date/as_of_date/date")
    as_of = parse_iso_or_slash_date(str(raw_date))
    rows = data.get("holdings")
    if rows is None:
        rows = data.get("positions")
    if rows is None:
        rows = data.get("stocks")
    if not isinstance(rows, list):
        rows = []
    holdings = []
    for h in rows:
        if not isinstance(h, dict):
            continue
        code = h.get("stock_code") or h.get("code") or h.get("ticker")
        if code is None or str(code).strip() == "":
            continue
        name = h.get("stock_name") or h.get("name") or ""
        shares = parse_int(h.get("shares", h.get("share", 0)))
        weight = parse_weight_pct(
            h.get("weight_pct", h.get("weight", h.get("weight_percent", 0)))
        )
        mv = h.get("market_value")
        if mv in (None, ""):
            mv = h.get("market_value_ntd", h.get("mv"))
        market_value = float(mv) if mv not in (None, "") else None
        holdings.append(
            _holding(
                code,
                name,
                shares,
                weight,
                policy,
                market_value=market_value,
            )
        )
    source_url = data.get("source_url") or data.get("official_url") or data.get("url")
    fetched_up = data.get("fetched_at") or data.get("fetched_at_upstream")
    return _snapshot(
        ticker=ticker,
        as_of_date=as_of,
        holdings=holdings,
        source="zdsetf",
        raw_path=path,
        meta={"source_url": source_url, "fetched_at_upstream": fetched_up},
    )


def _snapshot(
    *,
    ticker: str,
    as_of_date: date,
    holdings: list[dict[str, Any]],
    source: str,
    raw_path: Path,
    meta: dict[str, Any],
) -> dict[str, Any]:
    return {
        "etf_ticker": ticker,
        "as_of_date": as_of_date.isoformat(),
        "source": source,
        "fetched_at": now_taipei_iso(),
        "raw_file": raw_path.name,
        "raw_file_hash": file_sha256(raw_path),
        "holdings_count": len(holdings),
        "placeholder_count": sum(1 for h in holdings if h.get("is_placeholder")),
        "meta": meta,
        "holdings": holdings,
    }


def parse_excel_for_issuer(
    path: Path,
    ticker: str,
    issuer: str,
    cfg: AppConfig,
) -> dict[str, Any]:
    if issuer == "uni":
        return parse_uni_excel(path, ticker, cfg.placeholder)
    if issuer == "fh":
        return parse_fh_excel(path, ticker, cfg.placeholder)
    raise ValueError(f"unknown issuer: {issuer}")
